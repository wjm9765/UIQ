#!/usr/bin/env -S uv run python
from __future__ import annotations

import argparse
import html
import os
import re
import sys
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable

import numpy as np
import pandas as pd
import soundfile as sf
from datasets import load_dataset
from tqdm import tqdm

PROJECT_ROOT = Path(__file__).resolve().parents[1]
SRC_PATH = PROJECT_ROOT / "src"
if str(SRC_PATH) not in sys.path:
    sys.path.insert(0, str(SRC_PATH))

os.environ.setdefault("HF_HOME", str(PROJECT_ROOT / "input"))
os.environ.setdefault("HF_DATASETS_CACHE", str(PROJECT_ROOT / "input" / "datasets"))

from clap_eval.analysis.result_io import (  # noqa: E402
    MODEL_ORDER,
    ResultDataLoader,
    build_sample_id,
)
from clap_eval.config import Config  # noqa: E402

REQUIRED_MODELS = ["LAION", "M2D", "MGA", "MSCLAP"]
STATUS_NORMAL = "정상"
STATUS_ABNORMAL = "비정상"
MODEL_COMPARE_FIELDS = [
    "compare_audio_file",
    "compare_audio_src",
    "compare_split",
    "compare_dataset_index",
    "compare_youtube_id",
    "compare_sample_id",
    "compare_similarity",
]


@dataclass(frozen=True)
class OutputPaths:
    root: Path
    site: Path
    audio: Path
    excel: Path
    html: Path


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Build a human verification Excel file and static audio review site from raw CLAP outputs."
    )
    parser.add_argument("--config", default="config.yaml", help="Path to config.yaml")
    parser.add_argument("--input_dir", default=None, help="Directory with raw *_results.jsonl files")
    parser.add_argument("--output_dir", default="human_verify_output", help="Output directory")
    parser.add_argument("--reviewers", type=int, default=4, help="Number of reviewers")
    parser.add_argument("--per_status", type=int, default=25, help="Samples per reviewer for each status")
    parser.add_argument("--seed", type=int, default=42, help="Random seed for sampling")
    parser.add_argument(
        "--base_url",
        default="",
        help="Optional URL where the generated site/ directory will be hosted",
    )
    parser.add_argument(
        "--links_only",
        action="store_true",
        help="Only update Excel review_url hyperlinks in an existing output directory",
    )
    return parser.parse_args()


def make_output_paths(output_dir: str | Path) -> OutputPaths:
    root = Path(output_dir)
    site = root / "site"
    audio = site / "audio"
    root.mkdir(parents=True, exist_ok=True)
    site.mkdir(parents=True, exist_ok=True)
    audio.mkdir(parents=True, exist_ok=True)
    return OutputPaths(
        root=root,
        site=site,
        audio=audio,
        excel=root / "human_verify.xlsx",
        html=site / "index.html",
    )


def safe_filename(value: object, max_len: int = 90) -> str:
    text = str(value).strip().replace(os.sep, "_")
    text = re.sub(r"[^A-Za-z0-9._-]+", "_", text)
    text = re.sub(r"_+", "_", text).strip("._")
    return (text or "unknown")[:max_len]


def format_margin(value: object) -> float:
    if value is None or pd.isna(value):
        return float("nan")
    return float(value)


def build_review_url(base_url: str, review_id: str) -> str:
    if base_url:
        return f"{base_url.rstrip('/')}/index.html#{review_id}"
    return f"site/index.html#{review_id}"


def apply_excel_hyperlinks(worksheet) -> None:
    header_to_index = {
        str(cell.value): cell.column for cell in next(worksheet.iter_rows(min_row=1, max_row=1))
    }
    for header, column_index in header_to_index.items():
        if header != "review_url" and not header.endswith("audio_file"):
            continue
        for row_index in range(2, worksheet.max_row + 1):
            cell = worksheet.cell(row=row_index, column=column_index)
            if cell.value:
                cell.hyperlink = str(cell.value)
                cell.style = "Hyperlink"


def with_sample_ids(samples: pd.DataFrame) -> pd.DataFrame:
    samples = samples.copy()
    if "sample_id" not in samples.columns:
        samples["sample_id"] = (
            samples["split"].astype(str)
            + "|"
            + samples["index"].astype(str)
            + "|"
            + samples["youtube_id"].astype(str)
            + "|"
            + samples["start_time"].astype(str)
            + "|"
            + samples["label"].astype(str)
        )
    return samples


def load_raw_samples(input_dir: str | Path, model_order: Iterable[str]) -> pd.DataFrame:
    input_path = Path(input_dir)
    if input_path.is_file():
        raise ValueError(
            f"human_verify expects a raw result directory, not a cache file: {input_path}"
        )

    loader = ResultDataLoader(model_order=model_order)
    samples = loader.load(input_path)
    samples = samples[samples["model"].isin(model_order)].copy()
    if samples.empty:
        raise ValueError(f"No samples for required models were found under: {input_path}")
    return with_sample_ids(samples)


def compute_model_predictions(samples: pd.DataFrame, model_order: list[str]) -> pd.DataFrame:
    rows: list[dict] = []

    for model_name, model_df in tqdm(
        samples.groupby("model", dropna=False),
        desc="Prediction models",
        unit="model",
    ):
        if model_name not in model_order:
            continue

        model_df = model_df.copy()

        # Build a model-global unique text embedding table keyed by class label.
        text_source = model_df[model_df["text_embedding"].notna()]
        label_to_text: dict[str, np.ndarray] = {}
        for label, label_group in text_source.groupby("label", dropna=False):
            text_vec = label_group["text_embedding"].iloc[0]
            text_arr = np.asarray(text_vec, dtype=np.float32)
            text_arr = text_arr / (np.linalg.norm(text_arr) + 1e-8)
            label_to_text[str(label)] = text_arr

        if len(label_to_text) < 2:
            for row in model_df.to_dict("records"):
                rows.append(
                    {
                        "model": model_name,
                        "sample_id": str(row.get("sample_id", build_sample_id(row))),
                        "margin": np.nan,
                        "pred_label": "",
                    }
                )
            continue

        class_labels = sorted(label_to_text.keys())
        text_matrix = np.stack([label_to_text[label] for label in class_labels])
        label_to_idx = {label: idx for idx, label in enumerate(class_labels)}

        audio_matrix = np.stack(model_df["audio_embedding"].tolist())
        audio_matrix = audio_matrix / (np.linalg.norm(audio_matrix, axis=1, keepdims=True) + 1e-8)
        sim_matrix = audio_matrix @ text_matrix.T

        row_labels = model_df["label"].astype(str).tolist()
        pos_indices = np.array([label_to_idx.get(label, -1) for label in row_labels], dtype=np.int64)
        margins = np.full(len(model_df), np.nan, dtype=np.float32)
        pred_labels = np.array([""] * len(model_df), dtype=object)

        valid_mask = pos_indices >= 0
        if np.any(valid_mask):
            valid_rows = np.where(valid_mask)[0]
            valid_pos = pos_indices[valid_mask]

            valid_sims = sim_matrix[valid_mask].copy()
            positive = valid_sims[np.arange(len(valid_rows)), valid_pos]
            valid_sims[np.arange(len(valid_rows)), valid_pos] = -np.inf
            max_negative_indices = np.argmax(valid_sims, axis=1)
            max_negative = valid_sims[np.arange(len(valid_rows)), max_negative_indices]
            margins[valid_rows] = positive - max_negative

            # For strict top-1 correct cases, expose the GT label. Otherwise expose
            # the highest-scoring non-GT label so reviewers can inspect confusion.
            strict_correct = positive > max_negative
            display_pred_indices = np.where(strict_correct, valid_pos, max_negative_indices)
            pred_labels[valid_rows] = [class_labels[idx] for idx in display_pred_indices]

        for row, margin, pred_label in zip(
            model_df.to_dict("records"),
            margins,
            pred_labels,
        ):
            rows.append(
                {
                    "model": model_name,
                    "sample_id": str(row.get("sample_id", build_sample_id(row))),
                    "margin": float(margin),
                    "pred_label": str(pred_label),
                }
            )

    return pd.DataFrame.from_records(rows)


def build_candidate_table(samples: pd.DataFrame, model_order: list[str]) -> pd.DataFrame:
    print("[1/5] Computing per-model top-1 margins and predicted labels...")
    predictions = compute_model_predictions(samples, model_order=model_order)
    margin_matrix = predictions.pivot_table(
        index="sample_id",
        columns="model",
        values="margin",
        aggfunc="first",
    ).reindex(columns=model_order)
    pred_label_matrix = predictions.pivot_table(
        index="sample_id",
        columns="model",
        values="pred_label",
        aggfunc="first",
    ).reindex(columns=model_order)
    pred_label_matrix = pred_label_matrix.rename(
        columns={model_name: f"{model_name}_pred_label" for model_name in model_order}
    )

    metadata_cols = [
        "sample_id",
        "split",
        "index",
        "youtube_id",
        "start_time",
        "label",
        "meta_domain",
    ]
    metadata = (
        samples[metadata_cols]
        .drop_duplicates(subset=["sample_id"], keep="first")
        .set_index("sample_id")
    )

    candidates = (
        metadata.join(margin_matrix, how="inner")
        .join(pred_label_matrix, how="left")
        .reset_index()
    )
    candidates = candidates[candidates[model_order].notna().all(axis=1)].copy()
    candidates["status"] = np.select(
        [
            (candidates[model_order] > 0).all(axis=1),
            (candidates[model_order] <= 0).all(axis=1),
        ],
        [STATUS_NORMAL, STATUS_ABNORMAL],
        default="",
    )
    candidates = candidates[candidates["status"] != ""].copy()

    print(f"[1/5] Candidate samples with all four models present: {len(candidates):,}")
    print(f"      {STATUS_NORMAL}: {(candidates['status'] == STATUS_NORMAL).sum():,}")
    print(f"      {STATUS_ABNORMAL}: {(candidates['status'] == STATUS_ABNORMAL).sum():,}")
    return candidates


def choose_samples(
    candidates: pd.DataFrame,
    reviewers: int,
    per_status: int,
    seed: int,
) -> pd.DataFrame:
    total_per_status = reviewers * per_status
    rng = np.random.default_rng(seed)
    chosen_parts = []

    for status in [STATUS_NORMAL, STATUS_ABNORMAL]:
        pool = candidates[candidates["status"] == status]
        if len(pool) < total_per_status:
            raise ValueError(
                f"Not enough {status} samples: need {total_per_status}, found {len(pool)}"
            )
        chosen_index = rng.choice(pool.index.to_numpy(), size=total_per_status, replace=False)
        chosen = pool.loc[chosen_index].copy().reset_index(drop=True)
        chosen_parts.append(chosen)

    normal, abnormal = chosen_parts
    rows = []
    for reviewer in range(1, reviewers + 1):
        start = (reviewer - 1) * per_status
        end = start + per_status
        for status_code, status, subset in [
            ("N", STATUS_NORMAL, normal.iloc[start:end]),
            ("A", STATUS_ABNORMAL, abnormal.iloc[start:end]),
        ]:
            for local_idx, (_, row) in enumerate(subset.iterrows(), start=1):
                item = row.to_dict()
                item["reviewer"] = reviewer
                item["review_order"] = local_idx
                item["review_id"] = f"HV-R{reviewer}-{status_code}-{local_idx:03d}"
                rows.append(item)

    return pd.DataFrame.from_records(rows)


def add_comparison_audio_metadata(
    assignments: pd.DataFrame,
    samples: pd.DataFrame,
    model_order: list[str],
) -> pd.DataFrame:
    print("[2/5] Selecting nearest comparison audio for predicted wrong labels...")
    rows = assignments.to_dict("records")

    for row in rows:
        for model_name in model_order:
            prefix = f"{model_name}_"
            for field in MODEL_COMPARE_FIELDS:
                row[f"{prefix}{field}"] = ""

    for model_name in tqdm(model_order, desc="Comparison models", unit="model"):
        pred_col = f"{model_name}_pred_label"
        needed_rows = [
            row
            for row in rows
            if str(row.get(pred_col, "") or "")
            and str(row.get(pred_col, "") or "") != str(row["label"])
        ]
        if not needed_rows:
            continue

        needed_sample_ids = {str(row["sample_id"]) for row in needed_rows}
        needed_labels = {str(row[pred_col]) for row in needed_rows}
        model_df = samples[samples["model"] == model_name][
            [
                "sample_id",
                "split",
                "index",
                "youtube_id",
                "start_time",
                "label",
                "audio_embedding",
            ]
        ].copy()
        model_df["label"] = model_df["label"].astype(str)
        model_df["sample_id"] = model_df["sample_id"].astype(str)

        sample_vectors: dict[str, np.ndarray] = {}
        query_source = model_df[model_df["sample_id"].isin(needed_sample_ids)]
        for record in query_source.to_dict("records"):
            vector = np.asarray(record["audio_embedding"], dtype=np.float32)
            sample_vectors[str(record["sample_id"])] = vector / (np.linalg.norm(vector) + 1e-8)

        label_refs: dict[str, dict[str, object]] = {}
        ref_source = model_df[model_df["label"].isin(needed_labels)]
        for label, label_df in ref_source.groupby("label", dropna=False):
            label_rows = []
            vectors = []
            for record in label_df.to_dict("records"):
                vector = np.asarray(record["audio_embedding"], dtype=np.float32)
                vectors.append(vector / (np.linalg.norm(vector) + 1e-8))
                label_rows.append(
                    {
                        "sample_id": str(record["sample_id"]),
                        "split": str(record["split"]),
                        "index": as_optional_int(record["index"]),
                        "youtube_id": str(record["youtube_id"]),
                        "start_time": float(record["start_time"]),
                        "label": str(record["label"]),
                    }
                )

            if vectors:
                label_refs[str(label)] = {
                    "vectors": np.stack(vectors),
                    "rows": label_rows,
                }

        for row in needed_rows:
            prefix = f"{model_name}_"
            sample_id = str(row["sample_id"])
            pred_label = str(row[pred_col])
            query_vector = sample_vectors.get(sample_id)
            pred_refs = label_refs.get(pred_label)
            if query_vector is None or not pred_refs:
                continue

            vectors = pred_refs["vectors"]
            ref_rows = pred_refs["rows"]
            similarities = vectors @ query_vector
            best_idx = int(np.argmax(similarities))
            best_ref = ref_rows[best_idx]

            row[f"{prefix}compare_split"] = best_ref["split"]
            row[f"{prefix}compare_dataset_index"] = best_ref["index"]
            row[f"{prefix}compare_youtube_id"] = best_ref["youtube_id"]
            row[f"{prefix}compare_sample_id"] = best_ref["sample_id"]
            row[f"{prefix}compare_similarity"] = float(similarities[best_idx])

    return pd.DataFrame.from_records(rows)


def load_audio_split_cache(hf_repo: str, split: str, cache_dir: str, cache: dict[str, object]):
    if split not in cache:
        print(f"      Loading VGGSound split={split} from cache_dir={cache_dir}")
        cache[split] = load_dataset(
            hf_repo,
            split=split,
            cache_dir=cache_dir,
            trust_remote_code=True,
        )
    return cache[split]


def write_dataset_audio(
    *,
    hf_repo: str,
    split: str,
    dataset_index: int,
    cache_dir: str,
    split_cache: dict[str, object],
    site_audio_path: Path,
) -> None:
    dataset = load_audio_split_cache(hf_repo, split, cache_dir, split_cache)
    if dataset_index < 0 or dataset_index >= len(dataset):
        raise IndexError(
            f"Dataset index out of range: split={split}, index={dataset_index}, len={len(dataset)}"
        )

    if site_audio_path.exists():
        return

    audio = dataset[dataset_index]["audio"]
    sf.write(str(site_audio_path), audio["array"], audio["sampling_rate"])


def export_audio_files(
    assignments: pd.DataFrame,
    paths: OutputPaths,
    hf_repo: str,
    cache_dir: str,
    model_order: list[str],
) -> pd.DataFrame:
    print("[3/5] Exporting wav files for browser playback...")
    split_cache: dict[str, object] = {}
    rows = []

    for row in tqdm(assignments.to_dict("records"), desc="Writing wav", unit="file"):
        row = dict(row)
        reviewer = int(row["reviewer"])
        status = str(row["status"])
        status_dir = "normal" if status == STATUS_NORMAL else "abnormal"
        split = str(row["split"])
        dataset_index = int(row["index"])
        label = str(row["label"])

        filename = (
            f"{row['review_id']}__{safe_filename(split)}_{dataset_index}"
            f"__{safe_filename(label)}.wav"
        )
        relative_audio = Path("audio") / f"reviewer_{reviewer}" / status_dir / filename
        site_audio_path = paths.site / relative_audio
        site_audio_path.parent.mkdir(parents=True, exist_ok=True)

        write_dataset_audio(
            hf_repo=hf_repo,
            split=split,
            dataset_index=dataset_index,
            cache_dir=cache_dir,
            split_cache=split_cache,
            site_audio_path=site_audio_path,
        )

        row["audio_src"] = relative_audio.as_posix()
        row["audio_file"] = (Path("site") / relative_audio).as_posix()

        for model_name in model_order:
            compare_split = str(row.get(f"{model_name}_compare_split", "") or "")
            compare_index = as_optional_int(row.get(f"{model_name}_compare_dataset_index"))
            pred_label = str(row.get(f"{model_name}_pred_label", "") or "")
            if not compare_split or compare_index is None:
                row[f"{model_name}_compare_audio_src"] = ""
                row[f"{model_name}_compare_audio_file"] = ""
                continue

            compare_filename = (
                f"{safe_filename(model_name)}__{safe_filename(compare_split)}_{compare_index}"
                f"__{safe_filename(pred_label)}.wav"
            )
            compare_relative = Path("audio") / "comparison" / safe_filename(model_name) / compare_filename
            compare_site_path = paths.site / compare_relative
            compare_site_path.parent.mkdir(parents=True, exist_ok=True)

            write_dataset_audio(
                hf_repo=hf_repo,
                split=compare_split,
                dataset_index=compare_index,
                cache_dir=cache_dir,
                split_cache=split_cache,
                site_audio_path=compare_site_path,
            )

            row[f"{model_name}_compare_audio_src"] = compare_relative.as_posix()
            row[f"{model_name}_compare_audio_file"] = (Path("site") / compare_relative).as_posix()

        rows.append(row)

    return pd.DataFrame.from_records(rows)


def build_excel_table(assignments: pd.DataFrame, base_url: str) -> pd.DataFrame:
    rows = []
    for row in assignments.to_dict("records"):
        excel_row = {
            "담당자": int(row["reviewer"]),
            "검수ID": row["review_id"],
            "정상/비정상": row["status"],
            "카테고리": row["label"],
            "split": row["split"],
            "dataset_index": int(row["index"]),
            "review_url": build_review_url(base_url, row["review_id"]),
            "audio_file": row["audio_file"],
            "youtube_id": row["youtube_id"],
            "sample_id": row["sample_id"],
        }

        for model_name in REQUIRED_MODELS:
            excel_row.update(
                {
                    f"{model_name}_pred_label": row.get(f"{model_name}_pred_label", ""),
                    f"{model_name}_compare_audio_file": row.get(
                        f"{model_name}_compare_audio_file",
                        "",
                    ),
                    f"{model_name}_compare_split": row.get(
                        f"{model_name}_compare_split",
                        "",
                    ),
                    f"{model_name}_compare_dataset_index": row.get(
                        f"{model_name}_compare_dataset_index",
                        "",
                    ),
                    f"{model_name}_compare_youtube_id": row.get(
                        f"{model_name}_compare_youtube_id",
                        "",
                    ),
                    f"{model_name}_compare_sample_id": row.get(
                        f"{model_name}_compare_sample_id",
                        "",
                    ),
                    f"{model_name}_compare_similarity": row.get(
                        f"{model_name}_compare_similarity",
                        "",
                    ),
                    f"{model_name}_margin": format_margin(row[model_name]),
                }
            )

        excel_row.update({"사람판정": "", "메모": ""})
        rows.append(excel_row)
    return pd.DataFrame.from_records(rows)


def save_excel(excel_table: pd.DataFrame, output_path: Path) -> None:
    print("[4/5] Writing Excel workbook...")
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        excel_table.to_excel(writer, index=False, sheet_name="human_verify")
        worksheet = writer.sheets["human_verify"]
        apply_excel_hyperlinks(worksheet)
        for column_cells in worksheet.columns:
            header = str(column_cells[0].value)
            width = min(max(len(header) + 2, 12), 48)
            worksheet.column_dimensions[column_cells[0].column_letter].width = width
        worksheet.freeze_panes = "A2"


def update_existing_excel_links(excel_path: Path, base_url: str) -> None:
    if not excel_path.exists():
        raise FileNotFoundError(f"Excel file does not exist: {excel_path}")

    from openpyxl import load_workbook

    workbook = load_workbook(excel_path)
    worksheet = workbook["human_verify"]
    header_to_index = {cell.value: cell.column for cell in worksheet[1]}
    review_id_col = header_to_index.get("검수ID")
    review_url_col = header_to_index.get("review_url")
    if not review_id_col or not review_url_col:
        raise ValueError("Excel must contain 검수ID and review_url columns.")

    for row_index in range(2, worksheet.max_row + 1):
        review_id = worksheet.cell(row=row_index, column=review_id_col).value
        review_url = build_review_url(base_url, str(review_id))
        cell = worksheet.cell(row=row_index, column=review_url_col)
        cell.value = review_url

    apply_excel_hyperlinks(worksheet)
    workbook.save(excel_path)
    print(f"Updated Excel links: {excel_path}")


def html_cell(value: object) -> str:
    return html.escape("" if pd.isna(value) else str(value))


def as_optional_int(value: object) -> int | None:
    if value is None or pd.isna(value):
        return None
    if str(value).strip() == "":
        return None
    return int(value)


def render_margin(value: object) -> str:
    if value is None or pd.isna(value):
        return ""
    return f"{float(value):.4f}"


def render_similarity(value: object) -> str:
    if value is None or pd.isna(value) or str(value).strip() == "":
        return ""
    return f"{float(value):.4f}"


def render_model_cell(row: dict, model_name: str) -> str:
    pred_label = row.get(f"{model_name}_pred_label", "")
    compare_src = row.get(f"{model_name}_compare_audio_src", "")
    compare_split = row.get(f"{model_name}_compare_split", "")
    compare_index = row.get(f"{model_name}_compare_dataset_index", "")
    compare_similarity = render_similarity(row.get(f"{model_name}_compare_similarity", ""))
    margin = render_margin(row.get(model_name))

    if compare_src:
        audio_html = (
            f'<audio controls preload="none" src="{html_cell(compare_src)}"></audio>'
        )
        compare_meta = (
            f"{html_cell(compare_split)} #{html_cell(compare_index)}"
            + (f" · sim {html_cell(compare_similarity)}" if compare_similarity else "")
        )
    else:
        empty_label = "정답 라벨" if str(pred_label) == str(row.get("label", "")) else "비교 없음"
        audio_html = f'<span class="empty-audio">{html_cell(empty_label)}</span>'
        compare_meta = ""

    return f"""
      <td class="model-cell">
        <div class="pred-label">{html_cell(pred_label)}</div>
        {audio_html}
        <div class="submeta">{compare_meta}</div>
        <div class="submeta">margin {html_cell(margin)}</div>
      </td>
    """


def render_site(assignments: pd.DataFrame, output_path: Path) -> None:
    print("[5/5] Writing static review site...")
    generated_at = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")
    rows_html = []
    for row in assignments.to_dict("records"):
        rows_html.append(
            f"""
            <tr id="{html_cell(row['review_id'])}">
              <td>{html_cell(row['reviewer'])}</td>
              <td>{html_cell(row['review_id'])}</td>
              <td>{html_cell(row['status'])}</td>
              <td>{html_cell(row['label'])}</td>
              <td><audio controls preload="none" src="{html_cell(row['audio_src'])}"></audio></td>
              {render_model_cell(row, 'LAION')}
              {render_model_cell(row, 'M2D')}
              {render_model_cell(row, 'MGA')}
              {render_model_cell(row, 'MSCLAP')}
              <td>{html_cell(row['split'])}</td>
              <td>{html_cell(int(row['index']))}</td>
            </tr>
            """
        )

    html_doc = f"""<!doctype html>
<html lang="ko">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>Human Verification</title>
  <style>
    :root {{
      color-scheme: light;
      --bg: #f7f8fa;
      --panel: #ffffff;
      --text: #1f2933;
      --muted: #667085;
      --line: #d7dde5;
      --accent: #0f766e;
      --accent-soft: #e6f4f1;
    }}
    * {{ box-sizing: border-box; }}
    body {{
      margin: 0;
      background: var(--bg);
      color: var(--text);
      font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", sans-serif;
    }}
    header {{
      padding: 28px 32px 16px;
      border-bottom: 1px solid var(--line);
      background: var(--panel);
    }}
    h1 {{
      margin: 0;
      font-size: 24px;
      font-weight: 700;
      letter-spacing: 0;
    }}
    .meta {{
      margin-top: 8px;
      color: var(--muted);
      font-size: 14px;
    }}
    main {{ padding: 24px 32px 40px; }}
    .table-wrap {{
      overflow-x: auto;
      background: var(--panel);
      border: 1px solid var(--line);
      border-radius: 8px;
    }}
    table {{
      width: 100%;
      min-width: 1820px;
      border-collapse: collapse;
      font-size: 14px;
    }}
    th, td {{
      padding: 10px 12px;
      border-bottom: 1px solid var(--line);
      text-align: left;
      vertical-align: middle;
      white-space: nowrap;
    }}
    th {{
      position: sticky;
      top: 0;
      z-index: 1;
      background: var(--accent-soft);
      color: #0f3f3a;
      font-weight: 700;
    }}
    tr:target {{
      outline: 3px solid var(--accent);
      outline-offset: -3px;
      background: #f0fdfa;
    }}
    audio {{
      width: 260px;
      height: 36px;
    }}
    .model-cell {{
      min-width: 260px;
    }}
    .pred-label {{
      font-weight: 700;
      margin-bottom: 6px;
    }}
    .submeta {{
      margin-top: 4px;
      color: var(--muted);
      font-size: 12px;
    }}
    .empty-audio {{
      display: inline-flex;
      align-items: center;
      height: 36px;
      color: var(--muted);
      font-size: 13px;
    }}
    @media (max-width: 720px) {{
      header {{ padding: 20px 16px 12px; }}
      main {{ padding: 16px; }}
      h1 {{ font-size: 20px; }}
      audio {{ width: 220px; }}
    }}
  </style>
</head>
<body>
  <header>
    <h1>Human Verification</h1>
    <div class="meta">Rows: {len(assignments)} · Generated: {html_cell(generated_at)}</div>
  </header>
  <main>
    <div class="table-wrap">
      <table>
        <thead>
          <tr>
            <th>담당자</th>
            <th>검수ID</th>
            <th>정상/비정상</th>
            <th>카테고리</th>
            <th>원본 오디오</th>
            <th>LAION 예측/최근접 비교오디오</th>
            <th>M2D 예측/최근접 비교오디오</th>
            <th>MGA 예측/최근접 비교오디오</th>
            <th>MSCLAP 예측/최근접 비교오디오</th>
            <th>split</th>
            <th>dataset_index</th>
          </tr>
        </thead>
        <tbody>
          {''.join(rows_html)}
        </tbody>
      </table>
    </div>
  </main>
</body>
</html>
"""
    output_path.write_text(html_doc, encoding="utf-8")


def print_summary(excel_table: pd.DataFrame, paths: OutputPaths) -> None:
    print("\nDone.")
    print(f"Excel: {paths.excel}")
    print(f"Site:  {paths.html}")
    print("\nAssignment counts:")
    print(
        excel_table.groupby(["담당자", "정상/비정상"])
        .size()
        .unstack(fill_value=0)
        .to_string()
    )


def main() -> None:
    args = parse_args()
    config = Config.load(args.config)
    analysis_cfg = config.analysis
    dataset_cfg = config.dataset

    input_dir = args.input_dir or analysis_cfg.get("input_dir", "results/eval_outputs")
    output_paths = make_output_paths(args.output_dir)
    if args.links_only:
        update_existing_excel_links(output_paths.excel, args.base_url)
        return

    model_order = [m.upper() for m in MODEL_ORDER if m.upper() in REQUIRED_MODELS]

    required_total = args.reviewers * args.per_status
    print(f"Input raw results: {input_dir}")
    print(f"Output directory: {output_paths.root}")
    print(f"Required: {required_total} {STATUS_NORMAL} + {required_total} {STATUS_ABNORMAL}")

    samples = load_raw_samples(input_dir, model_order=model_order)
    candidates = build_candidate_table(samples, model_order=model_order)
    assignments = choose_samples(
        candidates,
        reviewers=args.reviewers,
        per_status=args.per_status,
        seed=args.seed,
    )
    assignments = add_comparison_audio_metadata(
        assignments,
        samples=samples,
        model_order=model_order,
    )

    assignments = export_audio_files(
        assignments,
        paths=output_paths,
        hf_repo=dataset_cfg.get("hf_repo", "txya900619/vggsound-16k"),
        cache_dir=dataset_cfg.get("cache_dir", "input"),
        model_order=model_order,
    )
    excel_table = build_excel_table(assignments, base_url=args.base_url)
    save_excel(excel_table, output_paths.excel)
    render_site(assignments, output_paths.html)
    print_summary(excel_table, output_paths)


if __name__ == "__main__":
    main()
